"""Fill SIP / QCP / checklist / SC list / FAI from a drawing JSON."""

from __future__ import annotations

import re
import shutil
import tempfile
import zipfile
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from sce_af.gdt_types import SPECIALS
from sce_af.paths import repo_root

TEMPLATE_DIR = "data/template"
SIP_TEMPLATE = "SIP.xlsx"
QCP_TEMPLATE = "QCP Template.xlsx"
CHECK_TEMPLATE = "SIP Checklist A1.xlsx"
SPEC_TEMPLATE = "SC list.xlsx"

DIM_HEADERS = ["序号", "标准来源", "测点号", "检查项目", "重要度", "名义值", "上限值", "下限值"]
GROUP_ORDER = ("重量", "性能", "尺寸")
STAGE_COLS = {"mold": 19}

_SIP_BODY_START = 11
_SIP_NOTE_MARK = "回溯全检要求"
_QCP_BODY_START = 11
_QCP_FOOTER_MARK = "表单编号"
_SPEC_BODY_START = 7
_SPEC_FOOTER_MARK = "备注"
_CHECK_BODY_START = 2
_CHECK_FOOTER_MARK = "表单编号"

_ILLEGAL = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

_FILE_KINDS = (
    ("检查作业标准", "write_sip"),
    ("控制计划", "write_qcp"),
    ("特殊特性清单", "write_spec"),
    ("检查清单", "write_checklist"),
    ("全尺寸FAI", "write_fai"),
)


def template_dir() -> Path:
    return repo_root() / TEMPLATE_DIR


def template_path(name: str) -> Path:
    path = template_dir() / name
    if not path.is_file():
        raise ValueError(f"缺少模板 {name}")
    return path


def header_text(header: dict, key: str) -> str:
    return str((header or {}).get(key) or "").strip()


def special_of(value: Any) -> str:
    text = str(value or "NA").strip().upper()
    return text if text in SPECIALS else "NA"


def check_item(item: dict) -> str:
    kind = str(item.get("type") or "").strip()
    desc = str(item.get("description") or "").strip()
    if kind and desc and not desc.startswith(kind):
        return f"{kind} {desc}"
    return desc or kind


def drawing_specials(bubbles: list[dict]) -> list[dict]:
    return [item for item in bubbles or [] if special_of(item.get("special")) in ("SC", "CC")]


def char_symbol(row: dict) -> str:
    if str(row.get("cc_no") or "").strip():
        return "CC"
    if str(row.get("sc_no") or "").strip():
        return "SC"
    return ""


def char_group(row: dict) -> str:
    name = str(row.get("name") or "")
    return "重量" if "重量" in name else "性能"


def export_stem(header: dict) -> str:
    parts = []
    for key in ("customer", "mold_no", "part_no", "revision"):
        text = header_text(header, key)
        if text:
            parts.append(_ILLEGAL.sub("-", text))
    return "-".join(parts) if parts else "导出"


def ascii_zip_name(stem: str, fallback: str) -> str:
    text = stem.encode("ascii", "ignore").decode("ascii").strip("-_. ")
    text = re.sub(r"-+", "-", text).strip("-")
    base = text or fallback or "docs"
    return f"{base}-docs.zip"


def fai_rows(doc: dict) -> list[dict]:
    rows = []
    for idx, item in enumerate(doc.get("bubbles") or [], start=1):
        rows.append(
            {
                "seq": idx,
                "source": "图纸",
                "point": str(item.get("num") if item.get("num") is not None else idx),
                "check": check_item(item),
                "grade": special_of(item.get("special")),
                "nominal": item.get("std_val"),
                "upper": item.get("max_val"),
                "lower": item.get("min_val"),
            }
        )
    return rows


def sip_lines(doc: dict) -> list[dict]:
    req = doc.get("requirements") or {}
    grouped: dict[str, list[dict]] = {name: [] for name in GROUP_ORDER}
    for row in req.get("characteristics") or []:
        group = char_group(row)
        grouped[group].append(
            {
                "group": group,
                "name": str(row.get("name") or "").strip(),
                "symbol": char_symbol(row),
                "spec": str(row.get("params") or "").strip(),
                "stages": [],
                "num": "",
            }
        )
    for item in drawing_specials(doc.get("bubbles") or []):
        num = str(item.get("num") if item.get("num") is not None else "")
        grouped["尺寸"].append(
            {
                "group": "尺寸",
                "name": num,
                "symbol": special_of(item.get("special")),
                "spec": str(item.get("description") or "").strip(),
                "stages": ["mold"],
                "num": num,
            }
        )
    out: list[dict] = []
    for group in GROUP_ORDER:
        out.extend(grouped[group])
    return out


def spec_rows(doc: dict) -> list[dict]:
    grouped: dict[str, list[dict]] = {name: [] for name in GROUP_ORDER}
    req = doc.get("requirements") or {}
    for row in req.get("characteristics") or []:
        symbol = char_symbol(row)
        if symbol not in ("SC", "CC"):
            continue
        group = char_group(row)
        grouped[group].append(
            {
                "category": group,
                "name": str(row.get("name") or "").strip(),
                "symbol": symbol,
                "desc": str(row.get("params") or "").strip(),
                "source": "图纸",
            }
        )
    for item in drawing_specials(doc.get("bubbles") or []):
        grouped["尺寸"].append(
            {
                "category": "尺寸",
                "name": check_item(item),
                "symbol": special_of(item.get("special")),
                "desc": str(item.get("description") or "").strip(),
                "source": "图纸",
            }
        )
    rows = []
    seq = 0
    for group in GROUP_ORDER:
        for row in grouped[group]:
            seq += 1
            rows.append({"seq": seq, **row})
    return rows


def qcp_rows(doc: dict) -> list[dict]:
    rows = []
    for item in drawing_specials(doc.get("bubbles") or []):
        rows.append(
            {
                "product": check_item(item),
                "symbol": special_of(item.get("special")),
                "spec": str(item.get("description") or "").strip(),
            }
        )
    return rows


def checklist_rows(doc: dict) -> list[dict]:
    dims = {str(row["point"]): row for row in fai_rows(doc)}
    out = []
    for seq, row in enumerate(sip_lines(doc), start=1):
        dim = dims.get(str(row.get("num") or row.get("name") or "")) if row.get("group") == "尺寸" else None
        name = row.get("name") or ""
        if dim:
            name = dim.get("check") or name
        out.append(
            {
                "seq": seq,
                "group": row.get("group") or "",
                "name": name,
                "symbol": row.get("symbol") or "",
                "spec": row.get("spec") or "",
                "lower": dim["lower"] if dim else None,
                "nominal": dim["nominal"] if dim else None,
                "upper": dim["upper"] if dim else None,
            }
        )
    return out


def _copy_row_style(ws, src_row: int, dst_row: int, max_col: int) -> None:
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
    if src_row in ws.row_dimensions:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def _set(ws, row: int, col: int, value: Any, align=None) -> None:
    if value is None or value == "":
        return
    cell = ws.cell(row, col)
    coord = cell.coordinate
    for rng in ws.merged_cells.ranges:
        if coord in rng:
            cell = ws.cell(rng.min_row, rng.min_col)
            break
    cell.value = value
    if align is not None:
        cell.alignment = align


def _find_row(ws, needle: str, col: int = 1) -> int:
    for row in range(1, ws.max_row + 1):
        if needle in str(ws.cell(row, col).value or ""):
            return row
    return 0


def _copy_row_merges(ws, src_row: int, dst_row: int) -> None:
    pairs = []
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row == src_row and rng.max_row == src_row and rng.min_col != rng.max_col:
            pairs.append((rng.min_col, rng.max_col))
    for min_col, max_col in pairs:
        ws.merge_cells(
            start_row=dst_row, start_column=min_col, end_row=dst_row, end_column=max_col
        )


def _merges_to_restore_after_insert(
    ws, insert_at: int, extra: int
) -> tuple[list[str], list[tuple[int, int, int, int]]]:
    stale: list[str] = []
    moved: list[tuple[int, int, int, int]] = []
    for rng in list(ws.merged_cells.ranges):
        min_r, max_r, min_c, max_c = rng.min_row, rng.max_row, rng.min_col, rng.max_col
        if max_r < insert_at:
            continue
        stale.append(str(rng))
        if min_r >= insert_at:
            min_r += extra
            max_r += extra
        else:
            max_r += extra
        if min_r != max_r or min_c != max_c:
            moved.append((min_r, max_r, min_c, max_c))
    return stale, moved


def _grow_before(ws, insert_at: int, extra: int, style_row: int, max_col: int, copy_merges: bool = False) -> None:
    if extra <= 0:
        return
    stale, moved = _merges_to_restore_after_insert(ws, insert_at, extra)
    for rng in stale:
        ws.unmerge_cells(rng)
    ws.insert_rows(insert_at, extra)
    for min_r, max_r, min_c, max_c in moved:
        ws.merge_cells(start_row=min_r, start_column=min_c, end_row=max_r, end_column=max_c)
    for i in range(extra):
        _copy_row_style(ws, style_row, insert_at + i, max_col)
        if copy_merges:
            _copy_row_merges(ws, style_row, insert_at + i)


def _unmerge_body(ws, start: int, end: int) -> None:
    stale = [
        str(rng)
        for rng in ws.merged_cells.ranges
        if rng.min_row >= start and rng.max_row < end
    ]
    for rng in stale:
        ws.unmerge_cells(rng)


def _sip_set_stages(ws, row: int, stages: list[str] | None) -> None:
    for col in STAGE_COLS.values():
        cell = ws.cell(row, col)
        if cell.value == "●":
            cell.value = None
    for name in stages or []:
        col = STAGE_COLS.get(str(name))
        if col:
            ws.cell(row, col).value = "●"


def write_fai(doc: dict, dest: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "尺寸检测报告"
    head_fill = PatternFill("solid", fgColor="D9E2F3")
    ws.append(DIM_HEADERS)
    for col in range(1, len(DIM_HEADERS) + 1):
        cell = ws.cell(1, col)
        cell.fill = head_fill
        cell.font = Font(bold=True)
        cell.alignment = _CENTER
    widths = (8, 10, 10, 34, 10, 12, 12, 12)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"
    for row in fai_rows(doc):
        ws.append(
            [
                row["seq"],
                row["source"],
                row["point"],
                row["check"],
                row["grade"],
                row["nominal"],
                row["upper"],
                row["lower"],
            ]
        )
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def write_sip(doc: dict, dest: Path) -> Path:
    shutil.copyfile(template_path(SIP_TEMPLATE), dest)
    wb = load_workbook(dest)
    ws = wb["检查标准书"]
    header = doc.get("header") or {}
    _set(ws, 3, 5, header_text(header, "part_name"), _LEFT)
    _set(ws, 3, 8, header_text(header, "mold_no"), _LEFT)
    _set(ws, 3, 14, header_text(header, "customer"), _LEFT)
    _set(ws, 4, 8, header_text(header, "part_no"), _LEFT)
    _set(ws, 4, 26, header_text(header, "revision"), _LEFT)

    lines = sip_lines(doc)
    note_row = _find_row(ws, _SIP_NOTE_MARK, col=2) or 35
    extra = max(0, len(lines) - max(0, note_row - _SIP_BODY_START))
    _grow_before(ws, note_row, extra, _SIP_BODY_START, 21, copy_merges=True)
    for offset, row in enumerate(lines):
        r = _SIP_BODY_START + offset
        _set(ws, r, 1, offset + 1, _CENTER)
        _set(ws, r, 2, row.get("group"), _LEFT)
        _set(ws, r, 3, row.get("name"), _LEFT)
        _set(ws, r, 4, row.get("symbol"), _CENTER)
        _set(ws, r, 6, row.get("spec"), _LEFT)
        if row.get("group") == "尺寸":
            _sip_set_stages(ws, r, row.get("stages"))
    wb.save(dest)
    return dest


def write_qcp(doc: dict, dest: Path) -> Path:
    shutil.copyfile(template_path(QCP_TEMPLATE), dest)
    wb = load_workbook(dest)
    ws = wb["控制计划"] if "控制计划" in wb.sheetnames else wb.active
    header = doc.get("header") or {}
    _set(ws, 2, 9, header_text(header, "customer"), _LEFT)
    _set(ws, 3, 6, header_text(header, "revision"), _LEFT)
    _set(ws, 4, 2, header_text(header, "part_name"), _LEFT)
    _set(ws, 6, 2, header_text(header, "mold_no"), _LEFT)
    _set(ws, 7, 2, header_text(header, "part_no"), _LEFT)

    lines = qcp_rows(doc)
    footer_row = _find_row(ws, _QCP_FOOTER_MARK) or ws.max_row + 1
    extra = max(0, len(lines) - max(0, footer_row - _QCP_BODY_START))
    _grow_before(ws, footer_row, extra, _QCP_BODY_START, 17)
    footer_row += extra
    _unmerge_body(ws, _QCP_BODY_START, footer_row)
    for r in range(_QCP_BODY_START, footer_row):
        for col in range(1, 18):
            ws.cell(r, col).value = None
    start = _QCP_BODY_START
    for offset, line in enumerate(lines):
        r = start + offset
        _set(ws, r, 7, line.get("product"), _LEFT)
        _set(ws, r, 9, line.get("symbol"), _CENTER)
        _set(ws, r, 10, line.get("spec"), _LEFT)
    if lines:
        end = start + len(lines) - 1
        _set(ws, start, 3, "P01", _CENTER)
        _set(ws, start, 4, "尺寸检验", _LEFT)
        if end > start:
            ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
            ws.merge_cells(start_row=start, start_column=4, end_row=end, end_column=4)
    wb.save(dest)
    return dest


def write_spec(doc: dict, dest: Path) -> Path:
    shutil.copyfile(template_path(SPEC_TEMPLATE), dest)
    wb = load_workbook(dest)
    ws = wb["Sheet1"]
    header = doc.get("header") or {}
    _set(ws, 2, 8, header_text(header, "revision"), _LEFT)
    _set(ws, 3, 4, header_text(header, "mold_no"), _LEFT)
    _set(ws, 4, 2, header_text(header, "part_name"), _LEFT)
    _set(ws, 4, 4, header_text(header, "part_no"), _LEFT)
    _set(ws, 4, 6, header_text(header, "customer"), _LEFT)

    rows = spec_rows(doc)
    footer = _find_row(ws, _SPEC_FOOTER_MARK) or ws.max_row + 1
    extra = max(0, len(rows) - max(0, footer - _SPEC_BODY_START))
    _grow_before(ws, footer, extra, _SPEC_BODY_START, 12, copy_merges=True)
    for offset, row in enumerate(rows):
        r = _SPEC_BODY_START + offset
        _set(ws, r, 1, row["seq"], _CENTER)
        _set(ws, r, 3, row.get("category"), _LEFT)
        _set(ws, r, 5, row["name"], _LEFT)
        _set(ws, r, 6, row["symbol"], _CENTER)
        _set(ws, r, 7, row["desc"], _LEFT)
        _set(ws, r, 9, row["source"], _LEFT)
    wb.save(dest)
    return dest


def write_checklist(doc: dict, dest: Path) -> Path:
    shutil.copyfile(template_path(CHECK_TEMPLATE), dest)
    wb = load_workbook(dest)
    ws = wb["检查清单表"] if "检查清单表" in wb.sheetnames else wb.active
    rows = checklist_rows(doc)
    footer = _find_row(ws, _CHECK_FOOTER_MARK) or 56
    extra = max(0, len(rows) - max(0, footer - _CHECK_BODY_START))
    _grow_before(ws, footer, extra, _CHECK_BODY_START, 18)
    for offset, row in enumerate(rows):
        r = _CHECK_BODY_START + offset
        _set(ws, r, 1, row["seq"], _CENTER)
        _set(ws, r, 2, row["group"], _LEFT)
        _set(ws, r, 3, row["name"], _LEFT)
        _set(ws, r, 5, row["symbol"], _CENTER)
        _set(ws, r, 7, row["spec"], _LEFT)
        _set(ws, r, 11, row["lower"], _CENTER)
        _set(ws, r, 12, row["nominal"], _CENTER)
        _set(ws, r, 13, row["upper"], _CENTER)
    wb.save(dest)
    return dest


_WRITERS = {
    "write_sip": write_sip,
    "write_qcp": write_qcp,
    "write_spec": write_spec,
    "write_checklist": write_checklist,
    "write_fai": write_fai,
}


def write_export_files(doc: dict, out_dir: Path) -> list[Path]:
    stem = export_stem(doc.get("header") or {})
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for label, key in _FILE_KINDS:
        dest = out_dir / f"{stem} {label}.xlsx"
        _WRITERS[key](doc, dest)
        paths.append(dest)
    return paths


def write_export_zip(doc: dict, zip_path: Path, fallback: str = "docs") -> str:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        files = write_export_files(doc, Path(tmp))
        with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for path in files:
                zf.write(path, arcname=path.name)
    return ascii_zip_name(export_stem(doc.get("header") or {}), fallback)
