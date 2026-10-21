from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import load_workbook

from sce_af.export_docs import (
    ascii_zip_name,
    export_stem,
    fai_rows,
    sip_lines,
    spec_rows,
    write_export_files,
    write_export_zip,
)

SAMPLE = {
    "header": {
        "filename": "LG-26554.pdf",
        "customer": "测试客户",
        "part_name": "测试品",
        "mold_no": "LG-26554",
        "part_no": "P000005766001",
        "drawing_no": "DWG-NOT-EXPORT",
        "revision": "A",
    },
    "bubbles": [
        {
            "num": "1",
            "special": "NA",
            "special_no": "",
            "type": "尺寸",
            "description": "4.7 0/-0.1",
            "page": 7,
            "nominal": 4.7,
            "upper": 0.0,
            "lower": -0.1,
            "min_val": 4.6,
            "std_val": 4.7,
            "max_val": 4.7,
        },
        {
            "num": "2",
            "special": "SC",
            "special_no": "13",
            "type": "位置度",
            "description": "位置度0.2",
            "page": 7,
            "nominal": 0.0,
            "upper": 0.2,
            "lower": 0.0,
            "min_val": 0.0,
            "std_val": 0.0,
            "max_val": 0.2,
        },
        {
            "num": "3",
            "special": "CC",
            "special_no": "4",
            "type": "面轮廓度",
            "description": "面轮廓度0.6",
            "page": 8,
            "nominal": 0.0,
            "upper": 0.6,
            "lower": 0.0,
            "min_val": 0.0,
            "std_val": 0.0,
            "max_val": 0.6,
        },
    ],
    "requirements": {
        "notes": [
            {
                "no": "1",
                "text": "条文不得出现XYZNOTE",
                "special": "NA",
                "special_no": "",
                "page": 4,
            }
        ],
        "characteristics": [
            {
                "cc_no": "",
                "sc_no": "01",
                "name": "产品总重量",
                "params": "0.087kg",
                "view": "VIEW-X",
                "zone": "ZONE-X",
                "location": "LOC-X",
                "page": 4,
            },
            {
                "cc_no": "02",
                "sc_no": "",
                "name": "漏电流",
                "params": "小于0.2mA",
                "view": "",
                "zone": "",
                "location": "",
                "page": 4,
            },
            {
                "cc_no": "",
                "sc_no": "",
                "name": "无标识外观",
                "params": "不得进清单",
                "view": "",
                "zone": "",
                "location": "",
                "page": 4,
            },
        ],
    },
    "issues": [{"code": "x", "message": "问题列表不得出现ABCISSUE"}],
}

FORBIDDEN = (
    "DWG-NOT-EXPORT",
    "条文不得出现XYZNOTE",
    "问题列表不得出现ABCISSUE",
    "VIEW-X",
    "ZONE-X",
    "LOC-X",
)


def _cell_texts(path: Path) -> list[str]:
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or cell.value == "":
                    continue
                out.append(str(cell.value))
    return out


def _by_suffix(paths: list[Path], needle: str) -> Path:
    for path in paths:
        if needle in path.name:
            return path
    raise AssertionError(f"missing {needle} in {[p.name for p in paths]}")


def test_row_models_split_json_blocks():
    assert len(fai_rows(SAMPLE)) == 3
    assert [row["grade"] for row in fai_rows(SAMPLE)] == ["NA", "SC", "CC"]
    lines = sip_lines(SAMPLE)
    assert [row["group"] for row in lines] == ["重量", "性能", "性能", "尺寸", "尺寸"]
    assert [row["name"] for row in lines] == ["产品总重量", "漏电流", "无标识外观", "2", "3"]
    specs = spec_rows(SAMPLE)
    assert [row["category"] for row in specs] == ["重量", "性能", "尺寸", "尺寸"]
    assert [row["symbol"] for row in specs] == ["SC", "CC", "SC", "CC"]
    assert [row["name"] for row in specs] == ["产品总重量", "漏电流", "位置度0.2", "面轮廓度0.6"]
    assert "无标识外观" not in [row["name"] for row in specs]
    assert export_stem(SAMPLE["header"]) == "测试客户-LG-26554-P000005766001-A"
    assert ascii_zip_name(export_stem(SAMPLE["header"]), "job") == "LG-26554-P000005766001-A-docs.zip"


def test_write_export_files_fill_templates(tmp_path):
    paths = write_export_files(SAMPLE, tmp_path)
    assert len(paths) == 5
    fai = load_workbook(_by_suffix(paths, "全尺寸FAI"))["尺寸检测报告"]
    assert fai.max_row == 4
    assert fai.cell(2, 3).value == "1"
    assert fai.cell(2, 5).value == "NA"
    assert fai.cell(2, 6).value == 4.7
    assert fai.cell(3, 5).value == "SC"
    assert fai.cell(4, 5).value == "CC"

    sip = load_workbook(_by_suffix(paths, "检查作业标准"))["检查标准书"]
    assert sip.cell(3, 5).value == "测试品"
    assert sip.cell(3, 8).value == "LG-26554"
    assert sip.cell(3, 14).value == "测试客户"
    assert sip.cell(4, 8).value == "P000005766001"
    assert sip.cell(4, 26).value == "A"
    assert sip.cell(11, 2).value == "重量"
    assert sip.cell(11, 3).value == "产品总重量"
    assert sip.cell(11, 4).value == "SC"
    assert sip.cell(11, 6).value == "0.087kg"
    assert sip.cell(12, 2).value == "性能"
    assert sip.cell(12, 3).value == "漏电流"
    assert sip.cell(13, 3).value == "无标识外观"
    assert sip.cell(14, 2).value == "尺寸"
    assert sip.cell(14, 3).value == "2"
    assert sip.cell(14, 19).value == "●"
    assert sip.cell(15, 3).value == "3"
    assert sip.cell(11, 5).value in (None, "")
    assert sip.cell(11, 14).value in (None, "")

    qcp = load_workbook(_by_suffix(paths, "控制计划"))["控制计划"]
    assert qcp.cell(2, 9).value == "测试客户"
    assert qcp.cell(6, 2).value == "LG-26554"
    assert qcp.cell(11, 3).value == "P01"
    assert qcp.cell(11, 4).value == "尺寸检验"
    assert qcp.cell(11, 9).value == "SC"
    assert qcp.cell(12, 9).value == "CC"
    assert "产品总重量" not in _cell_texts(_by_suffix(paths, "控制计划"))
    assert "漏电流" not in _cell_texts(_by_suffix(paths, "控制计划"))

    spec = load_workbook(_by_suffix(paths, "特殊特性清单"))["Sheet1"]
    assert spec.cell(7, 3).value == "重量"
    assert spec.cell(7, 5).value == "产品总重量"
    assert spec.cell(7, 6).value == "SC"
    assert spec.cell(7, 7).value == "0.087kg"
    assert spec.cell(8, 3).value == "性能"
    assert spec.cell(8, 5).value == "漏电流"
    assert spec.cell(8, 6).value == "CC"
    assert spec.cell(9, 3).value == "尺寸"
    assert spec.cell(9, 6).value == "SC"
    assert spec.cell(10, 6).value == "CC"
    assert spec.cell(7, 9).value == "图纸"
    assert spec.cell(7, 11).value in (None, "")
    spec_text = "\n".join(_cell_texts(_by_suffix(paths, "特殊特性清单")))
    assert "产品总重量" in spec_text
    assert "漏电流" in spec_text
    assert "无标识外观" not in spec_text
    assert "不得进清单" not in spec_text

    check = load_workbook(_by_suffix(paths, "检查清单"))["检查清单表"]
    assert check.cell(2, 2).value == "重量"
    assert check.cell(2, 3).value == "产品总重量"
    assert check.cell(3, 3).value == "漏电流"
    assert check.cell(4, 3).value == "无标识外观"
    assert check.cell(5, 2).value == "尺寸"
    assert "位置度0.2" in str(check.cell(5, 3).value)
    assert check.cell(5, 13).value == 0.2
    assert check.cell(2, 6).value in (None, "")

    for path in paths:
        blob = "\n".join(_cell_texts(path))
        for needle in FORBIDDEN:
            assert needle not in blob, f"{needle} in {path.name}"
        if "全尺寸FAI" not in path.name:
            assert "4.7 0/-0.1" not in blob
    fai_text = "\n".join(_cell_texts(_by_suffix(paths, "全尺寸FAI")))
    assert "NA" in fai_text
    assert sip.cell(16, 3).value in (None, "")


def test_na_stays_off_sip_and_qcp(tmp_path):
    sip = load_workbook(_by_suffix(write_export_files(SAMPLE, tmp_path), "检查作业标准"))["检查标准书"]
    names = [sip.cell(r, 3).value for r in range(11, 20)]
    assert "1" not in names
    qcp = load_workbook(_by_suffix(list(tmp_path.glob("*.xlsx")), "控制计划"))["控制计划"]
    products = [qcp.cell(r, 7).value for r in range(11, 20)]
    assert all(p is None or "4.7" not in str(p) for p in products)


def test_write_export_zip(tmp_path):
    zip_path = tmp_path / "docs.zip"
    name = write_export_zip(SAMPLE, zip_path, fallback="job1")
    assert name == "LG-26554-P000005766001-A-docs.zip"
    with zipfile.ZipFile(zip_path) as zf:
        names = zf.namelist()
    assert len(names) == 5
    assert any("检查作业标准" in n for n in names)
    assert any("控制计划" in n for n in names)
    assert any("检查清单" in n for n in names)
    assert any("特殊特性清单" in n for n in names)
    assert any("全尺寸FAI" in n for n in names)
